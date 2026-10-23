from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from split_data_files import main as split_data_files
from openpyxl.utils.cell import range_boundaries


WORKBOOK = Path.home() / "Downloads" / "\u9f99\u4e4b\u8c37" / "\u9f99\u4e4b\u8c37\u914d\u88c5\u6a21\u62df\u5668V2.12.xlsx"
OUTPUT = Path(__file__).resolve().parents[1] / "dragon_nest_builder" / "data" / "dragon_nest_data.json"


SHEETS = {
    "base": "\u6570\u636e_\u57fa\u7840",
    "enhance": "\u6570\u636e_\u5f3a\u5316",
    "suffix": "\u6570\u636e_\u540e\u7f00",
    "enchant": "\u6570\u636e_\u9644\u9b54",
    "jewelry": "\u6570\u636e_\u9970\u54c1",
    "set_bonus": "\u6570\u636e_\u5957\u88c5",
    "card_lookup": "lookup",
    "card_data": "\u6570\u636e_\u5361\u7247",
    "glyph_a": "\u6570\u636e_A\u7eb9\u7ae0",
    "glyph_a_extra": "\u6570\u636e_A\u989d\u5916",
    "glyph_s": "\u6570\u636e_S\u7eb9\u7ae0",
    "glyph_s_extra": "\u6570\u636e_S\u989d\u5916",
    "rune": "\u6570\u636e_\u77f3\u677f",
    "character": "\u6570\u636e_\u89d2\u8272\u57fa\u7840",
    "collection": "\u6570\u636e_\u6536\u96c6",
    "mastery": "\u6570\u636e_\u4e13\u7cbe",
    "title": "\u6570\u636e_\u79f0\u53f7",
    "fashion": "\u6570\u636e_\u65f6\u88c5",
    "conversion": "\u6570\u636e_\u5c5e\u6027\u8f6c\u6362\u6bd4\u4f8b",
    "job_map": "\u6570\u636e_\u804c\u4e1a\u8f6c\u6362\u6620\u5c04",
    "effective": "\u6570\u636e_\u7b49\u6548\u9762\u677f",
    "buff": "\u6570\u636e_\u6280\u80fdbuff",
    "auxiliary": "\u8f85\u52a9\u6570\u636e",
    "main": "\u603b\u914d\u88c5\u5668",
    "main2": "\u914d\u88c5\u56682",
    "total": "\u88c5\u5907+\u5361\u7247+\u7eb9\u7ae0+\u77f3\u677f\u603b\u5c5e\u6027",
}


def value(cell: Any) -> Any:
    if cell is None:
        return None
    return cell.value if hasattr(cell, "value") else cell


def rows_as_dicts(ws, max_col: int | None = None) -> list[dict[str, Any]]:
    headers = [value(c) for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=max_col))]
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True, max_col=max_col):
        item: dict[str, Any] = {}
        keep = False
        for i, header in enumerate(headers):
            if not header:
                continue
            cell_value = row[i] if i < len(row) else None
            if cell_value not in (None, ""):
                keep = True
            item[str(header)] = cell_value
        if keep:
            out.append(item)
    return out


def col_values(ws, col: int, start: int = 2) -> list[Any]:
    vals = []
    for row in range(start, ws.max_row + 1):
        v = ws.cell(row, col).value
        if v not in (None, ""):
            vals.append(v)
    return vals


def named_range_values(wb, name: str) -> list[Any]:
    defined = wb.defined_names.get(name)
    if not defined or not defined.attr_text or "!" not in defined.attr_text:
        return []
    sheet_name, ref = defined.attr_text.split("!", 1)
    sheet_name = sheet_name.strip("'")
    ref = ref.replace("$", "")
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    ws = wb[sheet_name]
    vals = []
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            v = ws.cell(row, col).value
            if v not in (None, ""):
                vals.append(v)
    return vals


def read_default_build(wb, sheet_name: str) -> dict[str, Any]:
    ws = wb[sheet_name]
    build: dict[str, Any] = {
        "first_job": ws["C3"].value,
        "second_job": ws["E3"].value,
        "crit_resist_value": ws["L37"].value or 0,
        "equipment": [],
        "jewelry": [],
        "glyph_a": [],
        "glyph_s": [],
        "normal_runes": [],
        "core_rune": {},
        "collection_level": ws["B56"].value or 0,
        "mastery": [],
        "title": ws["B69"].value,
        "fashion": [],
        "buffs": [],
        "card_grades": [],
        "card_bonus_levels": {},
    }
    for row in range(7, 14):
        build["equipment"].append({
            "slot": ws.cell(row, 1).value,
            "name": ws.cell(row, 2).value,
            "suffix": ws.cell(row, 3).value,
            "enhance": ws.cell(row, 4).value or 0,
            "enchants": [ws.cell(row, col).value for col in range(5, 9) if ws.cell(row, col).value],
        })
    for row in range(16, 20):
        build["jewelry"].append({
            "slot": ws.cell(row, 1).value,
            "name": ws.cell(row, 2).value,
            "variant": ws.cell(row, 3).value,
            "enchants": [ws.cell(row, col).value for col in range(4, 8) if ws.cell(row, col).value],
        })
    for row in range(23, 34):
        build["glyph_a"].append({
            "level": ws.cell(row, 1).value,
            "name": ws.cell(row, 2).value,
            "extra": ws.cell(row, 3).value,
        })
    for row in range(37, 41):
        build["glyph_s"].append({
            "level": ws.cell(row, 1).value,
            "name": ws.cell(row, 2).value,
            "base_attr1": ws.cell(row, 3).value,
            "rank1": ws.cell(row, 4).value,
            "base_attr2": ws.cell(row, 5).value,
            "rank2": ws.cell(row, 6).value,
            "extra": ws.cell(row, 7).value,
        })
    for row in range(44, 48):
        entries = []
        for col in range(3, 21, 3):
            entries.append({"attr": ws.cell(row, col).value, "rank": ws.cell(row, col + 1).value})
        build["normal_runes"].append({
            "type": ws.cell(row, 1).value,
            "level": ws.cell(row, 2).value,
            "entries": entries,
        })
    build["core_rune"] = {
        "level": ws["A51"].value,
        "conversion": ws["B51"].value,
        "ratio_rank": ws["C51"].value,
        "coeff_rank": ws["E51"].value,
        "enhance": ws["M51"].value,
    }
    for row in range(59, 67):
        build["mastery"].append({"name": ws.cell(row, 1).value, "level": ws.cell(row, 2).value or 0})
    for row in range(23, 32):
        build["fashion"].append({"slot": ws.cell(row, 6).value, "name": ws.cell(row, 7).value})
    for row in range(56, 66):
        build["buffs"].append({"skill": ws.cell(row, 9).value, "level": ws.cell(row, 10).value})
    for row in range(4, 95):
        build["card_grades"].append({
            "name": ws.cell(row, 25).value,
            "grade": ws.cell(row, 26).value or "\u65e0",
        })
    for col in range(27, 40):
        attr = ws.cell(3, col).value
        if attr:
            build["card_bonus_levels"][attr] = ws.cell(96, col).value or 0
    return build


def main() -> None:
    wb = load_workbook(WORKBOOK, data_only=False, read_only=False)
    wbc = load_workbook(WORKBOOK, data_only=True, read_only=False)
    data: dict[str, Any] = {
        "source": str(WORKBOOK),
        "tables": {},
        "lists": {},
        "defaults": {
            "build1": read_default_build(wbc, SHEETS["main"]),
            "build2": read_default_build(wbc, SHEETS["main2"]),
        },
    }
    for key, sheet in SHEETS.items():
        if key in {"main", "main2", "total"}:
            continue
        data["tables"][key] = rows_as_dicts(wb[sheet], max_col=9 if key == "buff" else None)

    total_ws = wb[SHEETS["total"]]
    data["attribute_order"] = [
        total_ws.cell(row, 1).value
        for row in range(2, total_ws.max_row + 1)
        if total_ws.cell(row, 1).value
    ]

    helper_glyph = wb["\u8f85\u52a9\u6570\u636e_\u7eb9\u7ae0"]
    helper_rune = wb["\u8f85\u52a9\u6570\u636e_\u77f3\u677f"]
    helper_title = wb["\u8f85\u52a9\u6570\u636e_\u79f0\u53f7"]
    helper_fashion = wb["\u8f85\u52a9\u6570\u636e_\u65f6\u88c5"]
    data["lists"] = {
        "first_jobs": named_range_values(wb, "JobClassList"),
        "glyph_a_levels": named_range_values(wb, "GlyphALevel"),
        "glyph_a_names": named_range_values(wb, "GlyphAName"),
        "glyph_a_extra": named_range_values(wb, "GlyphAPlus"),
        "glyph_s_levels": named_range_values(wb, "GlyphSLevel"),
        "glyph_s_names": named_range_values(wb, "GlyphSName"),
        "glyph_s_values": named_range_values(wb, "GlyphSValue"),
        "rune_levels": named_range_values(wb, "RuneLevel"),
        "rune_attack_attrs": named_range_values(wb, "RuneAttackAttr"),
        "rune_defense_attrs": named_range_values(wb, "RuneDefenseAttr"),
        "rune_ranks": named_range_values(wb, "RuneRank"),
        "core_levels": named_range_values(wb, "CoreRuneLevel"),
        "core_minus_attrs": named_range_values(wb, "CoreRuneMinusAttr"),
        "core_plus_attrs": named_range_values(wb, "CoreRunePlusAttr"),
        "core_ratio_ranks": named_range_values(wb, "CoreRuneRatio"),
        "titles": col_values(helper_title, 1, 2),
        "fashion": {
            helper_fashion.cell(1, col).value: col_values(helper_fashion, col, 2)
            for col in range(1, helper_fashion.max_column + 1)
            if helper_fashion.cell(1, col).value
        },
        "s_base_options": {
            helper_glyph.cell(1, col).value.replace("_\u57fa\u7840", ""): col_values(helper_glyph, col, 2)
            for col in range(10, 14)
            if helper_glyph.cell(1, col).value
        },
        "s_extra_options": {
            helper_glyph.cell(1, col).value.replace("_\u989d\u5916", ""): col_values(helper_glyph, col, 2)
            for col in range(15, 19)
            if helper_glyph.cell(1, col).value
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    split_data_files()
    print(OUTPUT)


if __name__ == "__main__":
    main()
