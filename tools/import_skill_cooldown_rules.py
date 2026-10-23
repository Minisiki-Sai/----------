import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = (
    ROOT
    / "outputs"
    / "cooldown_external_scenarios_20260718"
    / "技能减CD长期平均等效CD表_外部减CD情景.xlsx"
)
DEFAULT_OUTPUT = (
    ROOT
    / "dragon_nest_builder"
    / "data"
    / "modular"
    / "tables"
    / "skills"
    / "skill_cooldown_rules.json"
)

SOURCE_TYPES = {
    "固定秒数": "fixed_seconds",
    "重置冷却": "reset",
    "按目标CD百分比返还": "target_percent",
}

EFFECTIVE_SOURCE_LINKS = {
    ("时空领主", "九尾激光", "聚能激光"),
    ("时空领主", "光雨", "九尾激光"),
    ("时空领主", "激光破碎", "九尾激光"),
    ("时空领主", "能量爆炸", "九尾激光"),
}

def number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def compact_number(value: Any) -> int | float:
    parsed = number(value)
    return int(parsed) if parsed.is_integer() else parsed


def main() -> None:
    parser = argparse.ArgumentParser(description="从等效CD工作簿生成技能模拟器减CD规则。")
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    workbook = load_workbook(args.source, data_only=True, read_only=True)
    summary = workbook["等效CD汇总"]
    detail = workbook["减CD来源明细"]

    retained = {
        (str(row[0]).strip(), str(row[1]).strip())
        for row in summary.iter_rows(min_row=5, values_only=True)
        if row[0] and row[1]
    }
    rules: dict[str, dict[str, dict[str, Any]]] = {}

    for row in detail.iter_rows(min_row=5, values_only=True):
        job = str(row[0] or "").strip()
        target = str(row[1] or "").strip()
        if (job, target) not in retained:
            continue
        rule = rules.setdefault(job, {}).setdefault(
            target,
            {
                "base_cooldown_fallback": compact_number(row[2]),
                "passive_percent": 0,
                "self_refund_seconds": 0,
                "extra_resets": 0,
                "sources": [],
            },
        )
        source = str(row[3] or "").strip()
        source_type = str(row[5] or "").strip()
        if source_type == "常驻百分比":
            rule["passive_percent"] += number(row[9])
        elif source_type == "每周期自返还":
            rule["self_refund_seconds"] += number(row[6])
        elif source_type == "额外重置次数":
            rule["extra_resets"] += number(row[8])
        elif source_type in SOURCE_TYPES:
            source_rule: dict[str, Any] = {
                "skill": source,
                "type": SOURCE_TYPES[source_type],
                "source_cooldown_fallback": compact_number(row[4]),
            }
            if source_type == "固定秒数":
                source_rule["value"] = compact_number(row[6])
            elif source_type == "按目标CD百分比返还":
                source_rule["value"] = number(row[7])
            if source == "魔法之息触发":
                source_rule["affected_by_external_cdr"] = False
            if (job, target, source) in EFFECTIVE_SOURCE_LINKS:
                source_rule["use_effective_cooldown"] = True
            if (job, target, source) == ("时空领主", "九尾激光", "聚能激光"):
                # The worksheet's historical reset label means one stored use,
                # not a reset of every charge or merely the availability gate.
                source_rule.update(type="grant_charge", value=1)
            rule["sources"].append(source_rule)

    for job_rules in rules.values():
        for rule in job_rules.values():
            for field in ("passive_percent", "self_refund_seconds", "extra_resets"):
                rule[field] = compact_number(rule[field])

    payload = {
        "version": 1,
        "external_cdr_cap": 0.7,
        "description": "技能长期平均等效CD规则。基础CD优先读取当前技能数据，fallback仅在技能数据缺失时使用。",
        "jobs": dict(sorted(rules.items())),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(args.output)
    print(f"jobs={len(rules)}, skills={sum(len(items) for items in rules.values())}")


if __name__ == "__main__":
    main()
