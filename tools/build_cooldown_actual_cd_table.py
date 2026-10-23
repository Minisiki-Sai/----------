from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "exports" / "cooldown_reduction_skill_candidates.xlsx"
OUT = ROOT / "exports" / "cooldown_reduction_actual_cd_table.xlsx"


def main() -> None:
    source_wb = load_workbook(SRC, data_only=True)
    source_ws = source_wb.active
    rows = list(source_ws.iter_rows(min_row=2, values_only=True))

    wb = Workbook()
    ws = wb.active
    ws.title = "实际CD"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "减CD后实际技能CD表"
    ws["A1"].font = Font(bold=True, size=15, color="17324D")
    ws["A2"] = "说明：结果为减CD效果生效期间的面板CD；不计算覆盖率/循环期望。CD=0 的被动、EX触发或重置类条目标为不适用。"
    ws["A2"].font = Font(color="475569")
    ws.merge_cells("A1:J1")
    ws.merge_cells("A2:J2")

    assumptions = [
        ("鸡尾酒减CD", 0.12),
        ("精神爆发减CD", 0.40),
        ("双选减CD上限后", "=MIN(B5+B6,B9)"),
        ("时间加速减CD", 0.70),
        ("减CD上限", 0.70),
    ]
    ws["A4"] = "参数"
    ws["B4"] = "数值"
    for cell in ws["A4:B4"][0]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")
    for i, (label, value) in enumerate(assumptions, start=5):
        ws.cell(i, 1).value = label
        ws.cell(i, 2).value = value
        ws.cell(i, 2).number_format = "0%"

    headers = [
        "职业",
        "技能",
        "类型",
        "原始CD秒",
        "持续秒候选",
        "减CD说明/描述",
        "鸡尾酒后CD",
        "精神爆发后CD",
        "鸡尾酒+爆发后CD",
        "时间加速后CD",
    ]
    start_row = 11
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, row in enumerate(rows, start=start_row + 1):
        job, skill, typ, _level, _req, cd, dur, _cdr_value, desc, _file_name, _skill_id = row
        try:
            cd_num = float(cd or 0)
        except (TypeError, ValueError):
            cd_num = 0.0
        values = [job, skill, typ, cd_num if cd_num else 0, dur or "", desc or "", None, None, None, None]
        for col, value in enumerate(values, 1):
            ws.cell(idx, col).value = value
        if cd_num > 0:
            ws.cell(idx, 7).value = f"=ROUND(D{idx}*(1-$B$5),2)"
            ws.cell(idx, 8).value = f"=ROUND(D{idx}*(1-$B$6),2)"
            ws.cell(idx, 9).value = f"=ROUND(D{idx}*(1-$B$7),2)"
            ws.cell(idx, 10).value = f"=ROUND(D{idx}*(1-$B$8),2)"
        else:
            for col in range(7, 11):
                ws.cell(idx, col).value = "不适用"

    last_row = start_row + len(rows)
    table = Table(displayName="ActualCooldownTable", ref=f"A{start_row}:J{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    thin = Side(style="thin", color="CBD5E1")
    for row in ws.iter_rows(min_row=4, max_row=9, min_col=1, max_col=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=start_row + 1, max_row=last_row, min_col=1, max_col=10):
        for cell in row:
            cell.border = Border(bottom=Side(style="thin", color="E2E8F0"))
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column == 6))
    for col in range(4, 5):
        for cells in ws.iter_cols(min_col=col, max_col=col, min_row=start_row + 1, max_row=last_row):
            for cell in cells:
                cell.number_format = "0.00"
    for col in range(7, 11):
        for cells in ws.iter_cols(min_col=col, max_col=col, min_row=start_row + 1, max_row=last_row):
            for cell in cells:
                cell.number_format = "0.00"

    for i, width in enumerate([12, 18, 14, 10, 12, 56, 12, 13, 16, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row_index in range(start_row + 1, last_row + 1):
        ws.row_dimensions[row_index].height = 32
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:J{last_row}"

    quick = wb.create_sheet("简表")
    quick.sheet_view.showGridLines = False
    quick.append(["职业", "技能", "原始CD", "鸡尾酒", "爆发", "双选", "时间加速"])
    for cell in quick[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(horizontal="center")

    for row_index in range(start_row + 1, last_row + 1):
        base_cd = ws.cell(row_index, 4).value or 0
        if not isinstance(base_cd, (int, float)) or base_cd <= 0:
            continue
        quick.append(
            [
                ws.cell(row_index, 1).value,
                ws.cell(row_index, 2).value,
                f"='实际CD'!D{row_index}",
                f"='实际CD'!G{row_index}",
                f"='实际CD'!H{row_index}",
                f"='实际CD'!I{row_index}",
                f"='实际CD'!J{row_index}",
            ]
        )
    for col, width in enumerate([12, 18, 10, 10, 10, 10, 12], 1):
        quick.column_dimensions[get_column_letter(col)].width = width
    for row in quick.iter_rows(min_row=2):
        for cell in row[2:]:
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="right")
    quick.freeze_panes = "A2"
    if quick.max_row >= 2:
        quick_table = Table(displayName="CooldownQuickTable", ref=f"A1:G{quick.max_row}")
        quick_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        quick.add_table(quick_table)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)
    print(f"source rows={len(rows)}, quick active rows={quick.max_row - 1}")


if __name__ == "__main__":
    main()
