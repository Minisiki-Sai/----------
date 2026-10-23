from __future__ import annotations

import os
import re
import sys
from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QComboBox, QDoubleSpinBox

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from dragon_nest_builder.app import BuildEditor, SKILL_SIM_ROWS, as_number  # noqa: E402
from dragon_nest_builder.calculator import BuildCalculator  # noqa: E402
from dragon_nest_builder.data_loader import DataStore  # noqa: E402


OUTPUT = ROOT / "exports" / "skill_sim_dps_60L_all_jobs_v3_23.xlsx"


def combo_text(widget) -> str:
    return widget.currentText() if isinstance(widget, QComboBox) else ""


def item_text(table, row: int, col: int) -> str:
    item = table.item(row, col)
    return item.text() if item else ""


def clear_skill_buffs(editor: BuildEditor) -> None:
    for row in range(editor.buff_table.rowCount()):
        skill_widget = editor.buff_table.cellWidget(row, 0)
        if isinstance(skill_widget, QComboBox):
            editor._set_combo_current(skill_widget, "无")
        level_widget = editor.buff_table.cellWidget(row, 1)
        if isinstance(level_widget, QComboBox):
            editor._replace_combo_values(level_widget, ["无"], "无")


def set_all_cards_l(editor: BuildEditor) -> None:
    for row in range(editor.card_table.rowCount()):
        widget = editor.card_table.cellWidget(row, 1)
        if isinstance(widget, QComboBox):
            editor._set_combo_current(widget, "L")


def apply_attack_glyph_extra(editor: BuildEditor) -> None:
    profile = editor._job_profile()
    attack = profile.get("attack_attr") or "物攻"
    if isinstance(editor.auto_glyph_a_extra, QComboBox):
        editor._set_combo_current(editor.auto_glyph_a_extra, attack)
    editor.apply_auto_glyph_a_extra()


def calculate_skill_dps(editor: BuildEditor, store: DataStore, summary: dict[str, float]) -> tuple[float, float, list[dict[str, object]]]:
    weighted_panel = editor._skill_weighted_panel_profile()
    profile = editor._job_profile()
    preferred_type = None if weighted_panel else ("M" if profile.get("attack_attr") == "魔攻" else "P")
    total = 0.0
    rows: list[dict[str, object]] = []
    for row in range(SKILL_SIM_ROWS):
        skill = editor.combo_text(editor.skill_sim_table, row, 0)
        level = editor.combo_text(editor.skill_sim_table, row, 1)
        level_data = store.skill_damage_level_data(editor.second_job.currentText(), skill, level)
        cooldown_widget = editor.skill_sim_table.cellWidget(row, 3)
        cooldown = float(cooldown_widget.value()) if isinstance(cooldown_widget, QDoubleSpinBox) else 0.0
        multiplier_widget = editor.skill_sim_table.cellWidget(row, 4)
        multiplier = float(multiplier_widget.value()) if isinstance(multiplier_widget, QDoubleSpinBox) else 0.0
        element = store.skill_damage_element(editor.second_job.currentText(), skill)
        element_multiplier = editor._skill_element_display_multiplier(element, summary)
        dps_multiplier = editor._skill_sim_dps_multiplier(element, profile)
        effective = 0.0
        dps = 0.0
        dps_after = 0.0
        if level_data and cooldown > 0:
            effective = editor._skill_level_effective(
                level_data,
                summary,
                multiplier=multiplier,
                preferred_type=preferred_type,
                weighted_panel=weighted_panel,
            )
            dps = effective / cooldown * dps_multiplier
            dps_after = dps * element_multiplier
            total += dps
        rows.append(
            {
                "skill": skill,
                "level": level,
                "required_level": int(as_number(level_data[1])) if level_data else "",
                "cooldown": cooldown,
                "multiplier": multiplier,
                "element": element or "无",
                "effective_percent": effective,
                "element_multiplier": element_multiplier,
                "dps_percent": dps,
                "dps_after_element": dps_after,
            }
        )
    expected_panel = editor._skill_weighted_panel_value(summary, weighted_panel) if weighted_panel else as_number(summary.get("effective_panel"))
    return total, expected_panel * total / 100, rows


def fit_columns(sheet, max_width: int = 42) -> None:
    for col in range(1, sheet.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0.0
        for cell in sheet[letter]:
            text = "" if cell.value is None else str(cell.value)
            width = len(text) * 1.7 if re.search(r"[\u4e00-\u9fff]", text) else len(text)
            max_len = max(max_len, min(60, width))
        sheet.column_dimensions[letter].width = max(8, min(max_width, max_len + 2))


def style_workbook(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="DDEBFF")
    thin = Side(style="thin", color="D9E2F3")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"


def main() -> None:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    app = QApplication.instance() or QApplication([])
    store = DataStore.load()
    calculator = BuildCalculator(store)

    jobs: list[tuple[str, str]] = []
    for first in store.first_jobs():
        for second in store.second_jobs(first):
            if second and second != "无":
                jobs.append((first, second))

    summary_rows: list[dict[str, object]] = []
    detail_rows: list[dict[str, object]] = []
    errors: list[tuple[str, str, str]] = []

    for idx, (first, second) in enumerate(jobs, 1):
        try:
            build = deepcopy(store.defaults["build1"])
            build["first_job"] = first
            build["second_job"] = second
            build["buffs"] = []
            editor = BuildEditor(second, store, calculator, build)
            editor._export_skill_sim_page = editor._skill_sim_page()
            app.processEvents()

            editor.auto_equip("60L", "all")
            set_all_cards_l(editor)
            profile = editor._job_profile()
            if isinstance(editor.auto_glyph_a_extra, QComboBox):
                editor._set_combo_current(editor.auto_glyph_a_extra, profile.get("attack_attr") or "物攻")
            editor.auto_glyphs()
            apply_attack_glyph_extra(editor)
            editor.auto_misc()
            clear_skill_buffs(editor)
            editor.refresh()
            editor.auto_fill_skill_sim()
            editor.refresh()

            build = editor.collect_build()
            _result_rows, summary = calculator.calculate(build)
            editor.update_skill_sim_preview(summary)
            dps_total, expected_dps, skill_rows = calculate_skill_dps(editor, store, summary)
            element_attr = (profile.get("element_attr") or "").replace("%", "")

            summary_rows.append(
                {
                    "first_job": first,
                    "second_job": second,
                    "attack_type": profile.get("attack_attr", ""),
                    "element": element_attr or "无",
                    "effective_panel": as_number(summary.get("effective_panel")),
                    "physical_average": as_number(summary.get("physical_average")),
                    "magic_average": as_number(summary.get("magic_average")),
                    "final_percent": as_number(summary.get("final_percent")),
                    "crit_percent": as_number(summary.get("crit_percent")),
                    "crit_damage_percent": as_number(summary.get("crit_damage_percent")),
                    "dps_total": dps_total,
                    "expected_dps": expected_dps,
                }
            )

            for row, data in enumerate(skill_rows):
                detail_rows.append(
                    {
                        "first_job": first,
                        "second_job": second,
                        "row": row + 1,
                        "skill": data["skill"],
                        "level": data["level"],
                        "required_level": data["required_level"],
                        "cooldown": data["cooldown"],
                        "multiplier": data["multiplier"],
                        "element": data["element"],
                        "effective_percent": data["effective_percent"],
                        "element_multiplier": data["element_multiplier"],
                        "dps_percent": data["dps_percent"],
                        "dps_after_element": data["dps_after_element"],
                        "skill_value_text": item_text(editor.skill_sim_table, row, 5),
                        "effective_text": item_text(editor.skill_sim_table, row, 6),
                        "dps_text": item_text(editor.skill_sim_table, row, 7),
                    }
                )
            editor.deleteLater()
            print(f"[{idx}/{len(jobs)}] {first}-{second}: {dps_total:.3f}%/s")
        except Exception as exc:
            errors.append((first, second, repr(exc)))
            print(f"ERROR {first}-{second}: {exc!r}")

    summary_rows.sort(key=lambda row: float(row["dps_total"]), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "DPS排序"
    headers = [
        "排名",
        "一转职业",
        "二转职业",
        "攻击",
        "属性",
        "等效面板",
        "物攻均值",
        "魔攻均值",
        "最终%",
        "致命%",
        "爆伤%",
        "DPS总和(%/s)",
        "期望DPS/s",
        "建议系数_均值1",
        "建议系数_最高1",
        "建议系数_中位数1",
    ]
    ws.append(headers)
    for rank, row in enumerate(summary_rows, 1):
        ws.append(
            [
                rank,
                row["first_job"],
                row["second_job"],
                row["attack_type"],
                row["element"],
                row["effective_panel"],
                row["physical_average"],
                row["magic_average"],
                row["final_percent"],
                row["crit_percent"],
                row["crit_damage_percent"],
                row["dps_total"],
                row["expected_dps"],
                None,
                None,
                None,
            ]
        )
    last = len(summary_rows) + 1
    for row in range(2, last + 1):
        ws.cell(row, 14).value = f"=L{row}/AVERAGE($L$2:$L${last})"
        ws.cell(row, 15).value = f"=L{row}/MAX($L$2:$L${last})"
        ws.cell(row, 16).value = f"=L{row}/MEDIAN($L$2:$L${last})"

    ws2 = wb.create_sheet("技能模拟明细")
    ws2.append(
        [
            "一转职业",
            "二转职业",
            "序号",
            "技能",
            "等级",
            "所需等级",
            "冷却(s)",
            "倍率",
            "属性",
            "等效倍率%",
            "属性倍率",
            "DPS%/s",
            "属性后DPS%/s",
            "技能数值显示",
            "等效倍率显示",
            "DPS显示",
        ]
    )
    for row in detail_rows:
        ws2.append(
            [
                row["first_job"],
                row["second_job"],
                row["row"],
                row["skill"],
                row["level"],
                row["required_level"],
                row["cooldown"],
                row["multiplier"],
                row["element"],
                row["effective_percent"],
                row["element_multiplier"],
                row["dps_percent"],
                row["dps_after_element"],
                row["skill_value_text"],
                row["effective_text"],
                row["dps_text"],
            ]
        )

    ws3 = wb.create_sheet("配置和公式")
    for row in [
        ["导出条件", "每个职业：60L赞龙装备+60L赞龙首饰；所有卡片等级 L；一键纹章；A纹章额外属性尽量选择对应攻击；其他页一键全选后技能buff全部清空；最后执行一键模拟。"],
        ["DPS总和", "技能模拟中 10 个技能的“属性前DPS%/s”之和；属性后DPS仅作为明细参考。"],
        ["期望DPS/s", "普通职业=等效面板×DPS总和/100；十字军/圣骑士沿用当前模拟器的特殊技能面板口径。"],
        ["建议系数_均值1", "公式：本职业DPS总和 / 全职业DPS总和平均值。平均职业=1，越高说明技能表自身DPS越高。"],
        ["建议系数_最高1", "公式：本职业DPS总和 / 全职业DPS总和最大值。最高职业=1。"],
        ["建议系数_中位数1", "公式：本职业DPS总和 / 全职业DPS总和中位数。对极端值更不敏感。"],
        ["注意", "这些系数只用于观察技能模拟表的DPS差异，不代表实战循环、命中率、站桩时间、技能空转和团队环境。"],
    ]:
        ws3.append(row)

    if errors:
        ws4 = wb.create_sheet("错误")
        ws4.append(["一转职业", "二转职业", "错误"])
        for error in errors:
            ws4.append(list(error))

    style_workbook(wb)
    fit_columns(ws)
    fit_columns(ws2)
    fit_columns(ws3, max_width=80)
    for col in ["F", "G", "H", "M"]:
        ws.column_dimensions[col].width = 14
    for col in ["J", "K", "L", "M"]:
        ws2.column_dimensions[col].width = 14
    for col in ["N", "O", "P"]:
        ws2.column_dimensions[col].width = 36
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=16):
        for cell in row:
            cell.number_format = "0.000"
    for row in ws2.iter_rows(min_row=2, min_col=7, max_col=13):
        for cell in row:
            cell.number_format = "0.000"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"SAVED {OUTPUT}")
    print(f"jobs={len(summary_rows)} details={len(detail_rows)} errors={len(errors)}")


if __name__ == "__main__":
    main()
