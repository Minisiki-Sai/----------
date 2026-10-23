from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "龙之谷配装模拟器V4.01完整使用手册.md"
OUTPUT = ROOT / "docs" / "龙之谷配装模拟器V4.01完整使用手册.docx"

NAV_ITEMS = [
    ("01", "软件定位", "先理解四类结果和使用边界"),
    ("02", "运行准备", "发布版、源码版与 Npcap"),
    ("03", "十分钟上手", "从空白到第一套完整结果"),
    ("04–06", "配装与加点", "配装、对比、技能树与纹章"),
    ("07", "技能模拟", "DPS、DPA、爆发轴与长时间轴"),
    ("08", "配队模拟", "八人队伍、极限优化与机制轴"),
    ("09", "实战统计", "全场、Boss、时间轴、状态与评分"),
    ("10", "Hit 分析", "从记录学习 Hit 数与出伤时间"),
    ("11–12", "工作流与保存", "常见场景、备份和更新"),
    ("13–16", "排错与附录", "常见问题、术语和使用边界"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color="657786")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8.5, color="657786")


def add_inline(paragraph, text: str, *, color="243746", size=10.5) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.5, "0A6C82")
            run.font.highlight_color = None
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("243746")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "0B7285", 15, 8),
        ("Heading 2", 13, "155E75", 12, 6),
        ("Heading 3", 11.5, "1F4D78", 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Microsoft YaHei"
    styles["List Number"].font.name = "Microsoft YaHei"


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("DRAGON NEST BUILD SIMULATOR")
    set_run_font(run, "Calibri", 10, "0B7285", True)
    run.font.all_caps = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("龙之谷配装模拟器")
    set_run_font(run, size=28, color="102A43", bold=True)

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_after = Pt(20)
    run = version.add_run("V4.01")
    set_run_font(run, "Calibri", 34, "0B8793", True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    run = subtitle.add_run("完整使用手册")
    set_run_font(run, size=17, color="334E68", bold=True)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    band.columns[0].width = Inches(5.9)
    cell = band.cell(0, 0)
    set_cell_shading(cell, "E7F5F7")
    set_cell_margins(cell, 220, 300, 220, 300)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, "从第一次启动，到配装、加点、技能轴、八人配队与完整实战复盘", color="155E75", size=11.5)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    add_inline(p2, "写给从未使用、也未参与创作的新用户", color="486581", size=10)

    doc.add_paragraph().paragraph_format.space_after = Pt(60)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(meta, "文档日期  2026.08.07", color="829AB1", size=9.5)
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(author, "制作  帅气迷人机智幻", color="829AB1", size=9.5)

    doc.add_page_break()


def add_navigation(doc: Document) -> None:
    h = doc.add_heading("阅读导航", level=1)
    h.paragraph_format.space_before = Pt(0)
    p = doc.add_paragraph()
    add_inline(p, "第一次使用建议先阅读第 1–3 章并完成一套结果；遇到具体问题时，再按页面查阅后续章节。")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [0.75, 1.65, 4.7]
    for idx, text in enumerate(("章节", "主题", "你会学到")):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "0B7285")
        set_cell_margins(cell)
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_run_font(r, size=9.5, color="FFFFFF", bold=True)
    repeat_header(table.rows[0])
    for chapter, topic, goal in NAV_ITEMS:
        cells = table.add_row().cells
        values = (chapter, topic, goal)
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_margins(cell)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cell, "F2F8FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=9.2)
        keep_row_together(table.rows[-1])

    note = doc.add_table(rows=1, cols=1)
    note.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = note.cell(0, 0)
    set_cell_shading(cell, "FFF8E6")
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, "重要：配装面板、个人理论轴、八人配队轴和 .dpscap 实战记录是四种不同结果。比较数字前，先确认页面、Boss、时间范围和 Buff 环境相同。", color="7A4E00", size=9.5)

    doc.add_page_break()


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F5F8")
    set_cell_margins(cell, 120, 180, 120, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, "Consolas", 8.8, "243746")


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for idx in range(col_count):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "D9EEF2")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        add_inline(p, rows[0][idx] if idx < len(rows[0]) else "", color="134E5E", size=8.8)
        for run in p.runs:
            run.bold = True
    repeat_header(table.rows[0])
    keep_row_together(table.rows[0])

    for row_idx, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for idx in range(col_count):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 0:
                set_cell_shading(cell, "F6FAFB")
            p = cell.paragraphs[0]
            add_inline(p, values[idx] if idx < len(values) else "", size=8.7)
        keep_row_together(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def clean_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def build_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not line or line == "---":
            index += 1
            continue

        if line.startswith("# "):
            index += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            rows = [clean_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(clean_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, rows)
            continue

        if line.startswith(">"):
            quote_lines = []
            while index < len(lines) and (lines[index].startswith(">") or not lines[index].strip()):
                if lines[index].startswith(">"):
                    quote_lines.append(lines[index].lstrip("> ").rstrip())
                index += 1
            text = "  ".join(part for part in quote_lines if part)
            if text and "版本：" not in text and "文档日期：" not in text and "制作：" not in text:
                table = doc.add_table(rows=1, cols=1)
                cell = table.cell(0, 0)
                set_cell_shading(cell, "EAF6F8")
                set_cell_margins(cell, 130, 180, 130, 180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                add_inline(p, text, color="155E75", size=9.5)
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or number:
            if bullet:
                p = doc.add_paragraph(style="List Bullet")
                content = bullet.group(1)
            else:
                p = doc.add_paragraph()
                content = number.group(1)
                marker = re.match(r"^(\d+)\.", line).group(1)
                r = p.add_run(f"{marker}.  ")
                set_run_font(r, size=10.5, color="155E75", bold=True)
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, content)
            index += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        add_inline(p, line)
        index += 1


def add_headers_and_footers(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("龙之谷配装模拟器 V4.01  |  完整使用手册")
    set_run_font(r, size=8.5, color="5B7C8D", bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "B9DCE3")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_final_page(doc: Document) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    r = p.add_run("V4.01")
    set_run_font(r, "Calibri", 28, "0B8793", True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("让配装、时间轴、团队效果与实战记录互相验证")
    set_run_font(r, size=13, color="334E68", bold=True)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(22)
    add_inline(p3, "龙之谷配装模拟器 · 帅气迷人机智幻", color="829AB1", size=9.5)


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_headers_and_footers(doc)
    add_cover(doc)
    add_navigation(doc)
    build_body(doc, markdown)
    add_final_page(doc)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    core = doc.core_properties
    core.title = "龙之谷配装模拟器 V4.01 完整使用手册"
    core.subject = "配装、加点、技能模拟、八人配队与实战统计"
    core.author = "帅气迷人机智幻"
    core.keywords = "龙之谷, 配装模拟器, 技能模拟, 配队模拟, 实战统计"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
